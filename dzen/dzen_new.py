# -*- coding: utf-8 -*-
"""
Dzen视频数据抓取 - 改进版

使用方式:
  python dzen_new.py <window_id>
  python dzen_new.py <window_id> <channel_url>

输出:
  分析结果/dzen_videos_detail.xlsx
"""

import sys
import os
import time
import json
import openpyxl
import requests
from datetime import datetime

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(SCRIPT_DIR, "分析结果")
CONFIG_FILE = os.path.join(SCRIPT_DIR, "config.json")

sys.path.insert(0, SCRIPT_DIR)

from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from bit_api import openBrowser


def parse_number(num_str):
    if not num_str: return 0
    num_str = str(num_str).upper().strip().replace(',', '.').replace(' ', '')
    try:
        if 'ТЫС' in num_str: return int(float(num_str.replace('ТЫС', '').strip()) * 1000)
        if 'МЛН' in num_str: return int(float(num_str.replace('МЛН', '').strip()) * 1000000)
        if 'K' in num_str: return int(float(num_str.replace('K', '').strip()) * 1000)
        if 'M' in num_str: return int(float(num_str.replace('M', '').strip()) * 1000000)
        return int(float(num_str.replace('.', '').replace(',', '')))
    except: return 0


def check_bitbrowser():
    api_url = "http://127.0.0.1:54345"
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, 'r', encoding='utf-8') as f:
                config = json.load(f)
                api_url = config.get('bitbrowser_api', api_url)
        except: pass
    try:
        requests.get(api_url, timeout=5)
        return True
    except: return False


def connect_browser(window_id):
    res = openBrowser(window_id)
    if not res.get("success"):
        print(f"❌ 浏览器启动失败: {res.get('msg', '')}")
        return None
    options = webdriver.ChromeOptions()
    options.add_experimental_option("debuggerAddress", res["data"]["http"])
    options.add_argument("--disable-backgrounding-occluded-windows")
    driver = webdriver.Chrome(service=Service(res["data"]["driver"]), options=options)
    return driver


def scroll_and_get_videos(driver, max_scrolls=30):
    all_videos = {}
    no_new_count = 0
    print("滚动获取视频列表...")

    for i in range(max_scrolls):
        time.sleep(3)
        videos = driver.execute_script("""
            var videoList = [];
            // 修改：查找 /shorts/ 格式的链接（Dzen短视频真实格式）
            var links = document.querySelectorAll('a[href*="/shorts/"]');
            links.forEach(function(link) {
                // 从 /shorts/{ID} 提取视频ID
                var idMatch = link.href.match(/shorts\\/([a-zA-Z0-9]+)/);
                if (!idMatch || videoList.some(v => v.id === idMatch[1])) return;

                // 多种方式尝试获取标题
                var title = '';

                // 方式1: 从链接自身文本
                title = (link.innerText || '').trim();

                // 方式2: 从链接内的img alt属性
                if (!title) {
                    var img = link.querySelector('img');
                    if (img && img.alt) title = img.alt.trim();
                }

                // 方式3: 从父容器查找标题元素
                if (!title) {
                    var parent = link.closest('[class*="video-item"], [class*="card"], [class*="preview"]');
                    if (parent) {
                        var titleEl = parent.querySelector('[class*="title"], h2, h3, h4');
                        if (titleEl) title = (titleEl.innerText || '').trim();
                    }
                }

                // 方式4: 从aria-label属性
                if (!title) {
                    title = link.getAttribute('aria-label') || '';
                }

                videoList.push({
                    id: idMatch[1],
                    url: 'https://dzen.ru/shorts/' + idMatch[1],
                    title: title.substring(0, 80)
                });
            });
            return videoList;
        """)
        new_count = 0
        for v in videos:
            if v.get("id") and v["id"] not in all_videos:
                all_videos[v["id"]] = v
                new_count += 1
        print(f"  滚动{i+1}: 累计{len(all_videos)}个 (新增{new_count})")
        if new_count == 0:
            no_new_count += 1
            if no_new_count >= 10: break
        else: no_new_count = 0
        driver.execute_script("window.scrollBy(0, 1500);")
        time.sleep(3)
    return list(all_videos.values())


def get_video_detail(driver, video_url):
    try:
        driver.get(video_url)
        time.sleep(6)
        data = driver.execute_script("""
            var result = { views: 0, likes: 0, comments: 0, publishDate: '', title: '' };

            try {
                // 方法1: 从meta标签提取观看量
                var viewsMeta = document.querySelector('meta[property="ya:ovs:views_total"]');
                if (viewsMeta) {
                    result.views = parseInt(viewsMeta.getAttribute('content')) || 0;
                }

                // 方法2: 从meta标签提取发布时间
                var dateMeta = document.querySelector('meta[property="ya:ovs:upload_date"]');
                if (dateMeta) {
                    var dateStr = dateMeta.getAttribute('content') || '';
                    if (dateStr) {
                        result.publishDate = dateStr.split('T')[0];
                    }
                }

                // 方法3: 从JSON-LD提取点赞数和评论数
                var jsonLdScript = document.querySelector('script[type="application/ld+json"]');
                if (jsonLdScript) {
                    try {
                        var jsonData = JSON.parse(jsonLdScript.innerText);
                        if (jsonData.interactionStatistic) {
                            for (var stat of jsonData.interactionStatistic) {
                                if (stat.interactionType && stat.interactionType['@type'] === 'LikeAction') {
                                    result.likes = stat.userInteractionCount || 0;
                                }
                                if (stat.interactionType && stat.interactionType['@type'] === 'CommentAction') {
                                    result.comments = stat.userInteractionCount || 0;
                                }
                            }
                        }
                        if (jsonData.name) {
                            result.title = jsonData.name.substring(0, 80);
                        }
                    } catch(e) {}
                }

                // 方法4: 从H1元素提取标题（备用）
                if (!result.title) {
                    var titleEl = document.querySelector('h1[class*="short-detail__title"], h1');
                    if (titleEl && titleEl.innerText) {
                        result.title = titleEl.innerText.trim().substring(0, 80);
                    }
                }

                // 方法5: 如果meta标签没有观看量，尝试从页面元素提取
                if (result.views === 0) {
                    var pageText = document.body.innerText || '';
                    var viewsMatch = pageText.match(/([0-9,.]+[KkMmТЫС]*)\\s*(просмотров|просмотра)/i);
                    if (viewsMatch) {
                        var numStr = viewsMatch[1].replace(/[^0-9.,KkMmТЫС]/g, '');
                        if (numStr.indexOf('K') !== -1) result.views = parseFloat(numStr) * 1000;
                        else if (numStr.indexOf('M') !== -1) result.views = parseFloat(numStr) * 1000000;
                        else if (numStr.indexOf('ТЫС') !== -1) result.views = parseFloat(numStr.replace('ТЫС', '')) * 1000;
                        else result.views = parseInt(numStr.replace(/[^0-9]/g, '')) || 0;
                    }
                }

            } catch(e) {}
            return result;
        """)
        return data
    except Exception as e:
        print(f"    获取详情失败: {e}")
        return {"views": 0, "likes": 0, "comments": 0, "publishDate": "", "title": ""}


def save_excel(videos, output_file):
    if not videos: print("无数据"); return
    if not os.path.exists(OUTPUT_DIR): os.makedirs(OUTPUT_DIR)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dzen视频"
    headers = ["序号", "视频ID", "URL", "标题", "观看量", "点赞数", "评论数", "发布时间", "平台", "抓取时间"]
    for col, h in enumerate(headers, 1): ws.cell(1, col, h)
    for i, v in enumerate(videos, 2):
        ws.cell(i, 1, i-1)
        ws.cell(i, 2, v.get("id", ""))
        ws.cell(i, 3, v.get("url", ""))
        ws.cell(i, 4, v.get("title", ""))
        ws.cell(i, 5, v.get("views", 0))
        ws.cell(i, 6, v.get("likes", 0))
        ws.cell(i, 7, v.get("comments", 0))
        ws.cell(i, 8, v.get("publishDate", ""))
        ws.cell(i, 9, "Dzen")
        ws.cell(i, 10, datetime.now().strftime("%Y-%m-%d %H:%M"))
    ws.column_dimensions['C'].width = 50
    ws.column_dimensions['D'].width = 40
    wb.save(output_file)
    print(f"✅ 已保存: {output_file} ({len(videos)} 条)")


def main():
    print("=" * 60)
    print("Dzen视频数据抓取 - 改进版")
    print("=" * 60)

    if not check_bitbrowser():
        print("❌ BitBrowser 未运行，请先启动")
        return

    if len(sys.argv) < 2:
        print("\n用法: python dzen_new.py <window_id>")
        return

    window_id = sys.argv[1]
    channel_url = sys.argv[2] if len(sys.argv) > 2 else None

    # 从配置文件读取 channel_url（如果命令行未提供）
    if not channel_url and os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, 'r', encoding='utf-8') as f:
                config = json.load(f)
                channel_url = config.get('platforms', {}).get('dzen', {}).get('channel_url', '')
                if channel_url:
                    print(f"[配置] 使用频道URL: {channel_url}")
        except:
            pass

    print(f"\n连接浏览器 window_id: {window_id}")
    driver = connect_browser(window_id)
    if not driver: return

    try:
        handles = driver.window_handles
        for i in range(len(handles)-1, 0, -1):
            try: driver.switch_to.window(handles[i]); driver.close()
            except: pass
        driver.switch_to.window(driver.window_handles[0])

        if channel_url:
            print(f"\n访问频道: {channel_url}")
            driver.get(channel_url)
            time.sleep(5)

            # 自动导航到视频标签页
            import re
            current_url = driver.current_url
            match = re.search(r'(channel|id)\/([a-zA-Z0-9]+)', current_url)
            if match:
                # 先尝试 shorts（短视频），因为阀门视频多是短视频
                videos_url = f"https://dzen.ru/{match.group(1)}/{match.group(2)}?tab=shorts"
                print(f"访问短视频页面: {videos_url}")
                driver.get(videos_url)
        else:
            print("\n访问Dzen主页...")
            driver.get("https://dzen.ru/")
        time.sleep(6)

        # 先打印页面链接分析
        print("\n分析页面链接结构...")
        link_analysis = driver.execute_script("""
            var analysis = { allLinks: [], videoLinks: [], hrefPatterns: {} };
            var links = document.querySelectorAll('a');
            links.forEach(function(link) {
                var href = link.href || '';
                if (href.length > 10) {
                    analysis.allLinks.push(href.substring(0, 100));
                    // 分析 href 模式
                    if (href.indexOf('dzen.ru') !== -1) {
                        var pattern = href.replace(/https?:\/\/dzen\.ru/, '').split('?')[0];
                        if (pattern) {
                            analysis.hrefPatterns[pattern.substring(0, 50)] = (analysis.hrefPatterns[pattern.substring(0, 50)] || 0) + 1;
                        }
                    }
                    // 查找可能包含 video 的链接
                    if (href.match(/video|watch|stream|shorts|clip/i)) {
                        analysis.videoLinks.push(href);
                    }
                }
            });
            return analysis;
        """)

        print(f"总链接数: {len(link_analysis.get('allLinks', []))}")
        print(f"包含video关键词的链接: {len(link_analysis.get('videoLinks', []))}")
        print("链接模式统计:")
        for pattern, count in sorted(link_analysis.get('hrefPatterns', {}).items(), key=lambda x: -x[1])[:10]:
            print(f"  {pattern}: {count}次")
        if link_analysis.get('videoLinks'):
            print("\n视频相关链接示例:")
            for vlink in link_analysis.get('videoLinks', [])[:5]:
                print(f"  {vlink}")

        videos = scroll_and_get_videos(driver)
        print(f"\n获取到 {len(videos)} 个视频")

        if videos:
            print("\n获取详细数据...")
            for i, v in enumerate(videos):
                url = v.get("url", "")
                if not url: continue
                print(f"  {i+1}/{len(videos)}: {v.get('title', '')[:30]}")
                detail = get_video_detail(driver, url)
                v["views"] = detail.get("views", 0)
                v["likes"] = detail.get("likes", 0)
                v["comments"] = detail.get("comments", 0)
                v["publishDate"] = detail.get("publishDate", "")
                # 使用详情页提取的标题
                title_from_detail = detail.get("title", "")
                if title_from_detail:
                    v["title"] = title_from_detail
                    print(f"    观看:{v['views']} 点赞:{v['likes']} 标题:{title_from_detail[:40]}")
                else:
                    print(f"    观看:{v['views']} 点赞:{v['likes']} 标题:(空)")
                driver.back()
                time.sleep(2)

            output_file = os.path.join(OUTPUT_DIR, "dzen_videos_detail.xlsx")
            save_excel(videos, output_file)

            print("\n数据预览:")
            for v in videos[:10]:
                print(f"  {v.get('title', '')[:25]} | 观看:{v.get('views', 0)}")
        else:
            print("\n❌ 未获取到视频数据")

    except Exception as e:
        print(f"❌ 异常: {e}")
        import traceback
        traceback.print_exc()

    finally:
        try: driver.quit()
        except: pass

    print("=" * 60)


if __name__ == "__main__":
    main()