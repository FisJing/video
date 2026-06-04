# -*- coding: utf-8 -*-
"""
VK视频数据抓取 - 最终版

使用方式:
  python vk_final.py <window_id>
"""

import sys
import os
import time
import openpyxl
import requests
from datetime import datetime

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(SCRIPT_DIR, "分析结果")

from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from bit_api import openBrowser


def connect_browser(window_id):
    res = openBrowser(window_id)
    if not res.get("success"): return None
    options = webdriver.ChromeOptions()
    options.add_experimental_option("debuggerAddress", res["data"]["http"])
    driver = webdriver.Chrome(service=Service(res["data"]["driver"]), options=options)
    return driver


def navigate_to_clips(driver):
    # 直接导航到 clips 页面
    driver.get("https://cabinet.vkvideo.ru/dashboard")
    time.sleep(5)

    # 获取当前账户 ID
    current_url = driver.current_url
    account_id = ""
    if "@club" in current_url:
        account_id = current_url.split("@")[1].split("?")[0]

    # 直接打开 clips 页面
    clips_url = f"https://cabinet.vkvideo.ru/dashboard/@{account_id}?filterPreset=published&section=video_my_content&subsection=video_my_content_clips"
    driver.get(clips_url)
    time.sleep(8)
    print(f"当前URL: {driver.current_url}")


def extract_table_data(driver):
    print("提取表格数据...")
    videos = []

    # 滚动加载所有数据
    print("滚动加载所有数据...")

    # 先进行多次页面滚动
    for i in range(10):
        driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
        time.sleep(2)

    # 查找所有可滚动容器
    all_elements = driver.find_elements("css selector", "*")
    scrollable_elements = []
    for el in all_elements:
        try:
            scroll_height = int(el.get_attribute("scrollHeight") or 0)
            client_height = int(el.get_attribute("clientHeight") or 0)
            if scroll_height > client_height + 100:
                scrollable_elements.append(el)
        except: pass

    print(f"  找到 {len(scrollable_elements)} 个可滚动元素")

    # 滚动每个可滚动容器多次
    for container in scrollable_elements:
        for _ in range(10):
            try:
                driver.execute_script("arguments[0].scrollTop = arguments[0].scrollHeight", container)
                time.sleep(1)
            except: break

    # 再次页面滚动
    for i in range(5):
        driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
        time.sleep(2)

    # 检查行数
    table = driver.find_element("css selector", "table")
    rows = table.find_elements("css selector", "tr")
    print(f"表格总行数: {len(rows)}")

    # 提取数据
    table = driver.find_element("css selector", "table")
    rows = table.find_elements("css selector", "tr")
    print(f"表格总行数: {len(rows)}")

    for i, row in enumerate(rows[1:], 1):
        cells = row.find_elements("css selector", "td")
        if len(cells) < 5: continue

        video = {
            "title": cells[1].text.strip(),
            "date": cells[2].text.strip(),
            "status": cells[3].text.strip(),
            "views": int(cells[4].text.strip() or "0"),
            "likes": int(cells[5].text.strip() or "0"),
            "comments": int(cells[6].text.strip() or "0"),
            "shares": int(cells[7].text.strip() or "0"),
            "favorites": int(cells[8].text.strip() or "0"),
            "url": ""
        }
        links = cells[1].find_elements("tag name", "a")
        if links: video["url"] = links[0].get_attribute("href") or ""
        if video["title"]: videos.append(video)

    print(f"提取到: {len(videos)} 个视频")
    return videos


def save_excel(videos, output_file):
    if not videos: return
    if not os.path.exists(OUTPUT_DIR): os.makedirs(OUTPUT_DIR)
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = ["序号", "视频标题", "发布日期", "状态", "观看量", "点赞数", "评论数", "分享数", "收藏数", "URL", "平台", "抓取时间"]
    for col, h in enumerate(headers, 1): ws.cell(1, col, h)
    for i, v in enumerate(videos, 2):
        ws.cell(i, 1, i-1)
        ws.cell(i, 2, v.get("title", ""))
        ws.cell(i, 3, v.get("date", ""))
        ws.cell(i, 4, v.get("status", ""))
        ws.cell(i, 5, v.get("views", 0))
        ws.cell(i, 6, v.get("likes", 0))
        ws.cell(i, 7, v.get("comments", 0))
        ws.cell(i, 8, v.get("shares", 0))
        ws.cell(i, 9, v.get("favorites", 0))
        ws.cell(i, 10, v.get("url", ""))
        ws.cell(i, 11, "VK")
        ws.cell(i, 12, datetime.now().strftime("%Y-%m-%d %H:%M"))
    ws.column_dimensions['B'].width = 60
    wb.save(output_file)
    print(f"[OK] 已保存: {output_file} ({len(videos)} 条)")


def main():
    print("=" * 60)
    print("VK视频数据抓取 - 最终版")
    print("=" * 60)

    if len(sys.argv) < 2: print("用法: python vk_final.py <window_id>"); return

    driver = connect_browser(sys.argv[1])
    if not driver: return

    try:
        for i in range(len(driver.window_handles)-1, 0, -1):
            driver.switch_to.window(driver.window_handles[i]); driver.close()
        driver.switch_to.window(driver.window_handles[0])

        navigate_to_clips(driver)
        videos = extract_table_data(driver)

        if videos:
            save_excel(videos, os.path.join(OUTPUT_DIR, "vk_videos_final.xlsx"))
            print("\n数据预览:")
            for v in videos[:10]: print(f"  {v.get('title', '')[:40]} | 观看:{v.get('views', 0)}")
    except Exception as e:
        print(f"[FAIL] {e}")
    finally:
        try:
            driver.quit()
        except:
            pass
    print("=" * 60)


if __name__ == "__main__":
    main()